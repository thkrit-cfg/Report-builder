#!/usr/bin/env python3
"""Build a local daily budget bundle from the O2O budget workbook."""

from __future__ import annotations

import argparse
import json
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, default=Path("budget-data.js"))
    args = parser.parse_args()

    workbook = load_workbook(args.input, read_only=True, data_only=True)
    records = []
    for channel in ("QC", "PS", "TOL"):
        sheet = workbook[channel]
        days = [sheet.cell(3, column).value for column in range(6, 37)]
        values = [sheet.cell(5, column).value for column in range(6, 37)]
        for day, target in zip(days, values):
            if not isinstance(day, int) or target is None:
                continue
            records.append({
                "channel": channel,
                "date": date(2026, 8, day).isoformat(),
                "target": float(target),
            })

    args.output.write_text(
        "window.O2O_BUDGET_DATA=" + json.dumps(records, separators=(",", ":")) + ";\n",
        encoding="utf-8",
    )
    print(json.dumps({"rows": len(records), "output": str(args.output)}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
